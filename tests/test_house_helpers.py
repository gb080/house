import ast
from datetime import datetime
from pathlib import Path
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from zoneinfo import ZoneInfo


HOUSE_SOURCE = Path(__file__).parents[1] / "house.py"


def load_function(name, extra_globals=None):
    tree = ast.parse(HOUSE_SOURCE.read_text(encoding="utf-8"))
    function = next(node for node in tree.body if isinstance(node, ast.FunctionDef) and node.name == name)
    namespace = {
        "datetime": datetime,
        "manila_now": lambda: datetime.now(ZoneInfo("Asia/Manila")),
        "re": re,
        "Workbook": Workbook,
        "Font": Font,
        "PatternFill": PatternFill,
    }
    namespace.update(extra_globals or {})
    exec(compile(ast.Module(body=[function], type_ignores=[]), str(HOUSE_SOURCE), "exec"), namespace)
    return namespace[name], namespace


def test_parse_scanned_receipt_extracts_material_details():
    parse_scanned_receipt, _ = load_function("parse_scanned_receipt")

    result = parse_scanned_receipt("Cement\nQty: 4\nPHP 1,200.00\nDelivery 80")

    assert result == {"name": "Cement", "price": 300.0, "qty": 4, "delivery": 80.0}


def test_write_excel_keeps_payroll_expense_amount_under_net_column(tmp_path):
    month_key, month_globals = load_function("month_key")
    write_excel, write_globals = load_function("write_excel", {"month_key": month_key})
    write_globals["EXCEL_FILE"] = str(tmp_path / "ledger.xlsx")
    write_globals.update(month_globals)
    write_globals["EXCEL_FILE"] = str(tmp_path / "ledger.xlsx")

    write_excel({"payroll_expenses": [{"date": "Sep 11, 2026", "item": "Permit", "price": 500}]})
    workbook = load_workbook(tmp_path / "ledger.xlsx")
    sheet = workbook["Payroll"]

    assert sheet.max_column == 11
    assert sheet.cell(2, 5).value == "Permit"
    assert sheet.cell(2, 11).value == 500